"""构造合成规则表，只用于功能测试。不得写入真实体检报告目录。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

CANARY = "CANARY_TOKEN_ZX9"
MARKER_SHEET = "_synthetic_marker"
DATA_SHEET = "体检数据"

HEADERS = [
    "序号",
    "规则标识",
    "规则名称",
    "规则状态",
    "规则类型",
    "策略产品",
    "策略场景",
    "公司",
    "应用",
    "渠道/事件",
    "风险标签",
    "关联四级标签",
    "标签状态",
    "风险类型",
    "处置建议",
    "得分",
    "是否拉黑",
    "拉黑字段",
    "规则条件",
    "创建时间",
    "规则描述",
    "备注",
    "来源",
    "模型特征名称",
    "比较运算值",
    "阈值\\名单",
]


def _base(**overrides: str) -> dict[str, str]:
    row = {
        "序号": "1",
        "规则标识": "R_SINGLE",
        "规则名称": "合成规则甲",
        "规则状态": "已启用",
        "规则类型": "标准",
        "策略产品": "文本",
        "策略场景": "合成场景",
        "公司": "合成公司甲",
        "应用": "默认应用",
        "渠道/事件": "输入文本",
        "风险标签": "合成标签",
        "关联四级标签": "四级甲",
        "标签状态": "已启用",
        "风险类型": "其他",
        "处置建议": "拒绝",
        "得分": "900",
        "是否拉黑": "否",
        "拉黑字段": "",
        "规则条件": "(toxicity_score > 10)",
        "创建时间": "2026-01-01 00:00:00",
        "规则描述": "synthetic description",
        "备注": "synthetic",
        "来源": "合成来源甲",
        "模型特征名称": "toxicity_score",
        "比较运算值": ">",
        "阈值\\名单": "10",
        "_extra": "",
    }
    row.update(overrides)
    return row


def _values(row: dict[str, str]) -> list[str]:
    return [row[header] for header in HEADERS] + [row.get("_extra", "")]


def build_rows() -> list[dict[str, str]]:
    rows = [
        _base(),
        _base(序号="2", 规则标识="R_AND", 规则条件="(toxicity_score > 10)", 模型特征名称="toxicity_score", 比较运算值=">", **{"阈值\\名单": "10"}),
        _base(序号="2", 规则标识="R_AND", 规则条件="(insult_score >= 20)", 模型特征名称="insult_score", 比较运算值=">=", **{"阈值\\名单": "20"}),
        _base(序号="2", 规则标识="R_AND", 规则条件="(spam_score < 5)", 模型特征名称="spam_score", 比较运算值="<", **{"阈值\\名单": "5"}),
        _base(
            序号="3",
            规则标识="R_SCOPE",
            规则条件="(spam_score > 30)",
            模型特征名称="spam_score",
            比较运算值=">",
            **{"阈值\\名单": "30"},
            公司="合成公司甲",
            应用="默认应用",
            **{"渠道/事件": "输入文本"},
            来源="合成来源甲",
        ),
        _base(
            序号="4",
            规则标识="R_SCOPE",
            规则条件="(spam_score > 30)",
            模型特征名称="spam_score",
            比较运算值=">",
            **{"阈值\\名单": "30"},
            公司="合成公司乙",
            应用="B端应用",
            **{"渠道/事件": "输出文本"},
            来源="合成来源乙",
            规则类型="场景",
            策略场景="合成场景乙",
        ),
        _base(序号="5", 规则标识="R_BOTH", 公司="合成公司甲", 规则条件="(toxicity_score > 10)", 模型特征名称="toxicity_score", 比较运算值=">", **{"阈值\\名单": "10"}, 创建时间="2026-01-02 00:00:00"),
        _base(序号="5", 规则标识="R_BOTH", 公司="合成公司甲", 规则条件="(insult_score >= 20)", 模型特征名称="insult_score", 比较运算值=">=", **{"阈值\\名单": "20"}, 创建时间="2026-01-02 00:00:00"),
        _base(序号="6", 规则标识="R_BOTH", 公司="合成公司乙", 应用="B端应用", 来源="合成来源乙", 规则条件="(toxicity_score > 10)", 模型特征名称="toxicity_score", 比较运算值=">", **{"阈值\\名单": "10"}, 创建时间="2026-02-02 00:00:00"),
        _base(序号="6", 规则标识="R_BOTH", 公司="合成公司乙", 应用="B端应用", 来源="合成来源乙", 规则条件="(insult_score >= 20)", 模型特征名称="insult_score", 比较运算值=">=", **{"阈值\\名单": "20"}, 创建时间="2026-02-02 00:00:00"),
        _base(序号="7", 规则标识="R_CONFLICT", 规则条件="(ad_score > 40)", 模型特征名称="ad_score", 比较运算值=">", **{"阈值\\名单": "40"}, 处置建议="拒绝", 得分="900"),
        _base(序号="8", 规则标识="R_CONFLICT", 规则条件="(ad_score > 40)", 模型特征名称="ad_score", 比较运算值=">", **{"阈值\\名单": "40"}, 处置建议="审核", 得分="100"),
        _base(
            序号="9",
            规则标识="R_SHIFT",
            规则条件="(device_risk >= 5)",
            模型特征名称="account_hint",
            比较运算值="device_risk",
            **{"阈值\\名单": ">="},
            _extra="5",
        ),
        _base(
            序号="10",
            规则标识="R_BAD",
            规则条件=f"({CANARY})",
            模型特征名称="unused_feature",
            比较运算值=">",
            **{"阈值\\名单": "1"},
        ),
        _base(
            序号="11",
            规则标识="R_PART",
            规则条件="(score>=10 >= 1)",
            模型特征名称="score>=10",
            比较运算值=">=",
            **{"阈值\\名单": "1"},
        ),
        _base(
            序号="12",
            规则标识="R_NEQ",
            规则条件="(account_level neq guest)",
            模型特征名称="account_level",
            比较运算值="neq",
            **{"阈值\\名单": ""},
        ),
        _base(序号="13", 规则标识="R_DUP", 规则条件="(dup_score > 2)", 模型特征名称="dup_score", 比较运算值=">", **{"阈值\\名单": "2"}),
        _base(序号="13", 规则标识="R_DUP", 规则条件="(dup_score > 2)", 模型特征名称="dup_score", 比较运算值=">", **{"阈值\\名单": "2"}),
        _base(
            序号="14",
            规则标识="R_EMPTY",
            规则条件="",
            模型特征名称="",
            比较运算值="",
            **{"阈值\\名单": ""},
        ),
        _base(
            序号="15",
            规则标识="R_SUB",
            规则条件="(tag_set subsetp a,b,c)",
            模型特征名称="tag_set",
            比较运算值="subsetp",
            **{"阈值\\名单": "wrong"},
        ),
        _base(
            序号="16",
            规则标识="R_NULL",
            策略场景="",
            规则条件="(keyword_list contain hello)",
            模型特征名称="keyword_list",
            比较运算值="contain",
            **{"阈值\\名单": "hello"},
        ),
    ]
    return rows


def write_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = DATA_SHEET
    data.append(HEADERS + [None])
    for row in build_rows():
        data.append(_values(row))

    decoy = workbook.create_sheet("干扰表")
    decoy.append(["无关列"])
    for index in range(999):
        decoy.append([f"decoy-{index}"])

    marker = workbook.create_sheet(MARKER_SHEET)
    marker["A1"] = "healthcheck-synthetic"
    workbook.properties.subject = "healthcheck-synthetic"
    workbook.save(path)
    workbook.close()
