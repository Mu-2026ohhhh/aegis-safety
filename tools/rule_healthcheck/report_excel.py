"""把体检结果写成 Excel。调用方必须保证单元格里没有规则原文。"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from analyze import Analysis
from condition_parser import clean_excel_text


def write_report(analysis: Analysis, path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    sheets: list[tuple[str, pd.DataFrame]] = [
        ("00_说明", pd.DataFrame(analysis.cover, columns=["项目", "内容"])),
        ("01_字段概览", analysis.fields),
        ("02_分类分布", analysis.distributions),
        ("03_规则标识明细", analysis.rule_detail),
        ("04_多行与父行", analysis.structure),
        ("05_语义与版本差异", analysis.differences),
        ("06_条件解析", analysis.parse_stats),
        ("07_解析样例", analysis.samples),
        ("08_异常明细", analysis.anomalies),
        ("09_人工确认", analysis.review),
        ("10_后续影响", analysis.impacts),
    ]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name, frame in sheets:
            _safe_frame(frame).to_excel(writer, sheet_name=name, index=False)
        for worksheet in writer.book.worksheets:
            _format_sheet(worksheet, wide=worksheet.title in {"00_说明", "09_人工确认", "10_后续影响"})


def _safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    if frame is None or frame.empty and len(frame.columns) == 0:
        return pd.DataFrame()
    safe = frame.copy()
    for column in safe.columns:
        if safe[column].dtype == object:
            safe[column] = safe[column].map(lambda value: clean_excel_text(value) if isinstance(value, str) else value)
    return safe


def _format_sheet(worksheet, *, wide: bool) -> None:
    if worksheet.max_row < 1 or worksheet.max_column < 1:
        return
    worksheet.freeze_panes = "A2"
    if worksheet.max_row >= 1 and worksheet.max_column >= 1:
        worksheet.auto_filter.ref = worksheet.dimensions
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
    cap = 72 if wide else 42
    for index in range(1, worksheet.max_column + 1):
        header = worksheet.cell(1, index).value or ""
        width = min(max(len(str(header)) + 4, 14), cap)
        worksheet.column_dimensions[get_column_letter(index)].width = width
