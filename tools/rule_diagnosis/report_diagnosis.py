"""写出诊断 Excel 和 Markdown。不写回原始规则表。"""

from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

HEALTHCHECK_DIR = Path(__file__).resolve().parents[1] / "rule_healthcheck"
if str(HEALTHCHECK_DIR) not in sys.path:
    sys.path.insert(0, str(HEALTHCHECK_DIR))

from condition_parser import clean_excel_text  # noqa: E402
from compare import load_weights  # noqa: E402

SHEET_META: dict[str, list[tuple[str, str]]] = {
    "策略资产总览": [
        ("统计粒度", "同时给出原始行、规则标识、父行组、原子条件、规则与投放范围组合。"),
        ("纳入和排除", "纳入工作表全部数据行。解析不可用行计入总数，并单独列出。"),
        ("空值处理", "业务字段空单元格记为未配置；原文写成未配置则来源是原始业务值。"),
        ("异常计数口径", "空值率只数原始空单元格。"),
        ("适用限制", "未配置不是全渠道适用，也不是线上没有策略。"),
    ],
    "风险标签分布": [
        ("统计粒度", "原始行。同时给出规则标识数和父行组数，避免把一行当成一条规则。"),
        ("纳入和排除", "全部行。空标签显示为未配置。"),
        ("空值处理", "空值进入未配置，不并入其他标签。"),
        ("异常计数口径", "基数超过 30 时同时给 Top30 和完整明细。"),
        ("适用限制", "占比是行占比，不是线上流量占比。"),
    ],
    "风险类型分布": [
        ("统计粒度", "原始行，以及风险类型与处置建议的交叉。"),
        ("纳入和排除", "全部行，不把多个高基数字段再串成长组合。"),
        ("空值处理", "空的风险类型或处置建议记为未配置。"),
        ("异常计数口径", "交叉计数仍是行数。"),
        ("适用限制", "交叉分布不说明因果关系。"),
    ],
    "投放范围分布": [
        ("统计粒度", "公司、应用、渠道/事件、策略场景、来源分别计数；另给公司×渠道。"),
        ("纳入和排除", "模型特征和运算符只统计解析可用行，其余维度用全部行。"),
        ("空值处理", "空的投放维度记为未配置，不解释为全范围。"),
        ("异常计数口径", "高基数字段给 Top30 和完整明细。"),
        ("适用限制", "未配置的渠道不能当成所有渠道都已覆盖。"),
    ],
    "规则复杂度": [
        ("统计粒度", "每个规则标识一行，不把多行压成一条配置。"),
        ("纳入和排除", "包含解析不可用的行，并用列标明。"),
        ("空值处理", "列出出现过空值约定的业务字段。"),
        ("异常计数口径", "不一致指该标识下规范值多于一个，全部取值都保留。"),
        ("适用限制", "条件多或范围多不等于策略更危险。"),
    ],
    "完全重复候选": [
        ("统计粒度", "条件指纹、风险、处置、得分、标签状态和投放范围键都相同的一组行。"),
        ("纳入和排除", "单行组不输出。父行里不同的原子条件不会因为同属一个序号而进同一组。"),
        ("空值处理", "范围键使用未配置后的规范值。双方都空只表示这个维度都未配置。"),
        ("异常计数口径", "一组多行记一个候选。标记为完全重复候选。"),
        ("适用限制", "不是删除清单。"),
    ],
    "相同条件跨范围复用": [
        ("统计粒度", "一个条件指纹一行。范围不同或规则标识不同才纳入。"),
        ("纳入和排除", "同一规则、同一范围、只有一行的条件不纳入。"),
        ("空值处理", "有投放维度为未配置时单独标明，不称为全范围复用。"),
        ("异常计数口径", "处置不同、风险不同、未配置、解析问题可以同时为是，复用分类取优先展示项。"),
        ("适用限制", "跨范围复用可能是合理设计，不标成冗余。"),
    ],
    "相似规则候选": [
        ("统计粒度", "两个规则标识一对。相同条件指纹的对不在这里重复出现。"),
        ("纳入和排除", "只比较候选索引命中的对。解析不可用的特征不进入特征集合。"),
        ("空值处理", "双方未配置的维度贡献为 0，不增加排序参考。"),
        ("异常计数口径", "排序参考是加权和，只用于排序。超过 2000 对时截断。"),
        ("适用限制", "相似度不是重复概率、错误概率或风险概率。"),
    ],
    "同规则配置差异": [
        ("统计粒度", "一个规则标识一行，只保留处置、得分、标签或风险取值不同的标识。"),
        ("纳入和排除", "取值完全一致的规则不出现。"),
        ("空值处理", "未配置算一个规范值。空和原文未配置来源不同，但规范值都显示未配置。"),
        ("异常计数口径", "全部取值都列出。同一范围键或同一公司内的分歧单独标出。"),
        ("适用限制", "差异不是错误，也不要求立刻统一配置。"),
    ],
    "同特征阈值分布": [
        ("统计粒度", "解析可用的特征加运算符。名单另按指纹看跨规则引用。"),
        ("纳入和排除", "部分解析、字段不一致、疑似列错位全部排除，且不回填修正值。"),
        ("空值处理", "没有可靠阈值的行不进入本组。"),
        ("异常计数口径", "不同阈值个数按对外展示值去重。名单只显示指纹和长度。"),
        ("适用限制", "阈值多不等于阈值错误。"),
    ],
    "潜在处置冲突": [
        ("统计粒度", "同一规则的配置差异，以及相同条件指纹下处置不同的复用。"),
        ("纳入和排除", "没有处置差异的复用不进入本表。"),
        ("空值处理", "未配置的处置会和其他处置并列展示，不自动补成通过或拒绝。"),
        ("异常计数口径", "每条都标记为潜在冲突，待确认执行逻辑。"),
        ("适用限制", "不能据此声称线上已经冲突，也不知道优先级和聚合方式。"),
    ],
    "数据质量与排除记录": [
        ("统计粒度", "先给问题代码计数，再给不可用行或带质量代码的原始行。"),
        ("纳入和排除", "解析可用且没有质量代码的行不进明细。"),
        ("空值处理", "条件为空是技术质量问题，不写成业务未配置。"),
        ("异常计数口径", "一行可以有多个质量代码。计数按代码出现次数。"),
        ("适用限制", "质量问题说明分析受限，不直接等于线上误伤或漏放。"),
    ],
    "策略诊断清单": [
        ("统计粒度", "一条诊断覆盖一类问题，明细在前面的表。"),
        ("纳入和排除", "只收录本阶段规则命中的问题。没有效果数据的判断不写进去。"),
        ("空值处理", "诊断文本里的未配置沿用业务空值约定。"),
        ("异常计数口径", "优先级按对策略正确性的静态影响定性划分，不按出现次数自动升档。"),
        ("适用限制", "事实、假设、未知执行逻辑和建议动作分开写。当前状态都是待业务确认。"),
    ],
    "口径与限制说明": [
        ("统计粒度", "本表解释全部工作表和相似度权重。"),
        ("纳入和排除", "第二阶段没有独立的结构化成品。本阶段用体检解析器现算分析层。"),
        ("空值处理", "见各表。业务空值才叫未配置。"),
        ("异常计数口径", "见各表的口径项。"),
        ("适用限制", "不修改线上规则，不估计误杀率、漏放率或节省成本。"),
    ],
}


def write_workbook(path: Path, sheets: dict[str, pd.DataFrame]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    ordered = [name for name in SHEET_META if not name.startswith("_")]
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        for name in ordered:
            frame = sheets.get(name, pd.DataFrame())
            meta = pd.DataFrame(SHEET_META[name], columns=["口径项", "说明"])
            if name == "口径与限制说明":
                meta = pd.concat([meta, _weight_frame()], ignore_index=True)
            meta.to_excel(writer, sheet_name=name, index=False)
            start = len(meta) + 2
            _cells(frame).to_excel(writer, sheet_name=name, index=False, startrow=start)
        for worksheet in writer.book.worksheets:
            _format(worksheet)


def _weight_frame() -> pd.DataFrame:
    weights = load_weights()
    rows = [
        {"口径项": "第二阶段", "说明": "仓库里没有独立的第二阶段结构化文件。分析层在本次运行时由第一阶段解析器生成，质量问题只标记不修正。"},
        {"口径项": "相似度用途", "说明": weights["purpose"]},
        {"口径项": "候选索引", "说明": weights["candidate_index"]},
        {"口径项": "入选规则", "说明": weights["emit_rule"]},
        {"口径项": "双方未配置", "说明": "贡献固定为 0，不因为两边都空而更相似。"},
    ]
    for name, spec in weights["dimensions"].items():
        rows.append({"口径项": f"权重 {name}={spec['weight']}", "说明": spec["method"]})
    return pd.DataFrame(rows)


def _cells(frame: pd.DataFrame) -> pd.DataFrame:
    if frame is None or len(frame.columns) == 0:
        return pd.DataFrame()
    safe = frame.copy()
    for column in safe.columns:
        if safe[column].dtype == object:
            safe[column] = safe[column].map(lambda value: clean_excel_text(value) if isinstance(value, str) else value)
    return safe


def _format(worksheet) -> None:
    if worksheet.max_row < 1:
        return
    fill = PatternFill("solid", fgColor="1F4E79")
    font = Font(color="FFFFFF", bold=True)
    for cell in worksheet[1]:
        cell.fill = fill
        cell.font = font
    for index in range(1, min(worksheet.max_column, 12) + 1):
        worksheet.column_dimensions[get_column_letter(index)].width = 28


def render_markdown(source_name: str, sheet_name: str, data_kind: str, counts: dict[str, int], diagnosis: pd.DataFrame) -> str:
    kind = "合成测试" if data_kind == "synthetic" else "真实数据"
    lines = [
        f"# 策略诊断摘要（{kind}）",
        "",
        f"- 源文件名：{source_name}",
        f"- 工作表：{sheet_name}",
        "- 第二阶段没有单独的结构化产物。本次使用规则体检里的条件解析器生成分析层，原始值保留，解析失败和字段不一致没有被改写。",
        "- 业务字段空值记为“未配置”。这不等于全范围适用，也不等于线上没有覆盖。",
        "",
        "## 资产规模",
        "",
        f"- 原始行 {counts['原始行数']}，规则标识 {counts['规则标识数']}，父行组 {counts['父行组数']}。",
        f"- 原子条件 {counts['原子条件数']}，规则与投放范围组合 {counts['规则与投放范围组合数']}。",
        f"- 解析不可用行 {counts['解析不可用行数']}。这些行不进入阈值分布和特征相似度。",
        "",
        "## 诊断数量",
        "",
        f"- 完全重复候选 {counts['完全重复候选组数']} 组。",
        f"- 相同条件跨范围复用 {counts['跨范围复用条件数']} 个条件指纹。",
        f"- 相似规则候选 {counts['相似规则候选对数']} 对。排序参考不是概率。",
        f"- 同规则配置差异 {counts['同规则配置差异数']} 个规则标识。",
        f"- 潜在处置冲突 {counts['潜在处置冲突条数']} 条，状态都是“待确认执行逻辑”。",
        f"- 诊断清单 {counts['诊断条目数']} 条：P0 {counts['P0']}，P1 {counts['P1']}，P2 {counts['P2']}，P3 {counts['P3']}。",
        "",
        "## 优先人工核查",
        "",
    ]
    if len(diagnosis):
        focus = diagnosis[diagnosis["处理优先级"].isin(["P0", "P1"])]
        if len(focus) == 0:
            lines.append("本次没有 P0 或 P1 条目。")
        for record in focus.to_dict("records"):
            lines.append(f"### {record['问题编号']} {record['处理优先级']} {record['问题类型']}")
            lines.append("")
            lines.append(f"- 事实：{record['发现的事实']}")
            lines.append(f"- 理由：{record['优先级理由']}")
            lines.append(f"- 建议：{record['建议核查动作']}")
            lines.append("")
    lines.extend(
        [
            "## 当前不能成立的结论",
            "",
            "- 不能说线上已经发生规则冲突。表里没有优先级、聚合、白名单和后置处置。",
            "- 不能把跨范围复用或完全重复候选直接当成应删除的冗余。",
            "- 不能把未配置解释成全渠道适用或确认无覆盖。",
            "- 不能给出误杀率、漏放率、效果提升或成本节省。没有线上命中和人工复核数据。",
            "",
            "## 下一阶段需要的数据",
            "",
            "- 规则命中日志：内容、命中规则标识、处置结果、时间。",
            "- 执行顺序或聚合说明：多规则同时命中时如何选处置。",
            "- 人工复核或申诉样本：用来区分配置差异和真实误伤、漏放。",
            "",
        ]
    )
    return "\n".join(lines)
