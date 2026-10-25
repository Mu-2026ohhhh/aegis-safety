"""合成规则表，只用于诊断测试。"""

from __future__ import annotations

from pathlib import Path

import openpyxl

CANARY = "CANARY_TOKEN_ZX9"
MARKER_SHEET = "_synthetic_marker"
DATA_SHEET = "诊断样本"

HEADERS = [
    "序号",
    "规则标识",
    "规则名称",
    "规则状态",
    "规则类型",
    "策略产品",
    "策略场景",
    "公司",
    "应用",
    "渠道/事件",
    "风险标签",
    "关联四级标签",
    "标签状态",
    "风险类型",
    "处置建议",
    "得分",
    "是否拉黑",
    "拉黑字段",
    "规则条件",
    "创建时间",
    "规则描述",
    "备注",
    "来源",
    "模型特征名称",
    "比较运算值",
    "阈值\\名单",
]


def _base(**overrides: str) -> dict[str, str]:
    row = {
        "序号": "1",
        "规则标识": "R_SINGLE",
        "规则名称": "合成规则",
        "规则状态": "已启用",
        "规则类型": "标准",
        "策略产品": "文本",
        "策略场景": "合成场景",
        "公司": "合成公司甲",
        "应用": "默认应用",
        "渠道/事件": "输入文本",
        "风险标签": "合成标签",
        "关联四级标签": "四级甲",
        "标签状态": "已启用",
        "风险类型": "其他",
        "处置建议": "拒绝",
        "得分": "900",
        "是否拉黑": "否",
        "拉黑字段": "",
        "规则条件": "(toxicity_score > 10)",
        "创建时间": "2026-01-01 00:00:00",
        "规则描述": "synthetic",
        "备注": "synthetic",
        "来源": "合成来源",
        "模型特征名称": "toxicity_score",
        "比较运算值": ">",
        "阈值\\名单": "10",
        "_extra": "",
    }
    row.update(overrides)
    return row


def build_rows() -> list[dict[str, str]]:
    threshold = "阈值\\名单"
    channel = "渠道/事件"
    return [
        _base(序号="1", 规则标识="R_PARENT", 规则条件="(toxicity_score > 10)", 模型特征名称="toxicity_score", 比较运算值=">", **{threshold: "10"}),
        _base(序号="1", 规则标识="R_PARENT", 规则条件="(insult_score >= 20)", 模型特征名称="insult_score", 比较运算值=">=", **{threshold: "20"}, 规则名称="父行拆分"),
        _base(序号="2", 规则标识="R_REUSE", 公司="合成公司甲", 规则条件="(spam_score > 30)", 模型特征名称="spam_score", 比较运算值=">", **{threshold: "30"}),
        _base(序号="3", 规则标识="R_REUSE", 公司="合成公司乙", 规则条件="(spam_score > 30)", 模型特征名称="spam_score", 比较运算值=">", **{threshold: "30"}),
        _base(序号="4", 规则标识="R_DUP_A", 规则条件="(dup_score > 2)", 模型特征名称="dup_score", 比较运算值=">", **{threshold: "2"}, 得分="800", 风险类型="广告"),
        _base(序号="5", 规则标识="R_DUP_B", 规则条件="(dup_score > 2)", 模型特征名称="dup_score", 比较运算值=">", **{threshold: "2"}, 得分="800", 风险类型="广告"),
        _base(序号="6", 规则标识="R_CONFLICT", 公司="合成公司甲", 规则条件="(ad_score > 40)", 模型特征名称="ad_score", 比较运算值=">", **{threshold: "40"}, 处置建议="拒绝", 得分="900"),
        _base(序号="7", 规则标识="R_CONFLICT", 公司="合成公司乙", 规则条件="(ad_score > 40)", 模型特征名称="ad_score", 比较运算值=">", **{threshold: "40"}, 处置建议="审核", 得分="100"),
        _base(序号="8", 规则标识="R_SAMECO", 规则条件="(same_scope > 3)", 模型特征名称="same_scope", 比较运算值=">", **{threshold: "3"}, 处置建议="拒绝"),
        _base(序号="9", 规则标识="R_SAMECO", 规则条件="(same_scope > 3)", 模型特征名称="same_scope", 比较运算值=">", **{threshold: "3"}, 处置建议="通过", 备注="另一条备注"),
        _base(
            序号="10",
            规则标识="R_SHIFT",
            规则条件="(device_risk >= 5)",
            模型特征名称="account_hint",
            比较运算值="device_risk",
            **{threshold: ">="},
            _extra="5",
        ),
        _base(
            序号="11",
            规则标识="R_PART",
            规则条件="(score>=10 >= 1)",
            模型特征名称="score>=10",
            比较运算值=">=",
            **{threshold: "1"},
        ),
        _base(
            序号="12",
            规则标识="R_SIM_A",
            规则名称="合成辱骂阈值甲",
            风险类型="辱骂",
            规则条件="(insult_score >= 20)",
            模型特征名称="insult_score",
            比较运算值=">=",
            **{threshold: "20", channel: ""},
        ),
        _base(
            序号="13",
            规则标识="R_SIM_B",
            规则名称="合成辱骂阈值乙",
            风险类型="辱骂",
            规则条件="(insult_score >= 30)",
            模型特征名称="insult_score",
            比较运算值=">=",
            **{threshold: "30", channel: ""},
        ),
        _base(序号="14", 规则标识="R_ACT1", 规则条件="(cross_score > 8)", 模型特征名称="cross_score", 比较运算值=">", **{threshold: "8"}, 处置建议="拒绝"),
        _base(序号="15", 规则标识="R_ACT2", 规则条件="(cross_score > 8)", 模型特征名称="cross_score", 比较运算值=">", **{threshold: "8"}, 处置建议="审核"),
        _base(
            序号="16",
            规则标识="R_LIST1",
            风险类型="广告",
            公司="合成公司甲",
            规则条件="(tag_set subsetp a,b,c)",
            模型特征名称="tag_set",
            比较运算值="subsetp",
            **{threshold: "a,b,c"},
        ),
        _base(
            序号="17",
            规则标识="R_LIST2",
            风险类型="违禁",
            公司="合成公司乙",
            规则条件="(tag_set subsetp a,b,c)",
            模型特征名称="tag_set",
            比较运算值="subsetp",
            **{threshold: "a,b,c"},
        ),
        _base(序号="18", 规则标识="R_BAD", 规则条件=f"({CANARY})", 模型特征名称="unused_feature", 比较运算值=">", **{threshold: "1"}),
        _base(序号="19", 规则标识="R_BLANK", **{channel: ""}, 规则条件="(blank_channel > 1)", 模型特征名称="blank_channel", 比较运算值=">", **{threshold: "1"}),
        _base(序号="20", 规则标识="R_LITERAL", **{channel: "未配置"}, 规则条件="(literal_channel > 1)", 模型特征名称="literal_channel", 比较运算值=">", **{threshold: "1"}),
    ]


def write_workbook(path: Path) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    workbook = openpyxl.Workbook()
    data = workbook.active
    data.title = DATA_SHEET
    data.append(HEADERS + [None])
    for row in build_rows():
        data.append([row[header] for header in HEADERS] + [row.get("_extra", "")])
    decoy = workbook.create_sheet("干扰表")
    decoy.append(["无关列"])
    decoy.append(["不要读入"])
    marker = workbook.create_sheet(MARKER_SHEET)
    marker["A1"] = "healthcheck-synthetic"
    workbook.properties.subject = "healthcheck-synthetic"
    workbook.save(path)
    workbook.close()
